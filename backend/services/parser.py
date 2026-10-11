import os
import fitz  # PyMuPDF
from docx import Document
from pptx import Presentation
import pandas as pd
from PIL import Image

def parse_pdf(file_path: str):
    """Parses a PDF file page by page, returning text and metadata."""
    results = []
    doc = fitz.open(file_path)
    for page_num in range(len(doc)):
        page = doc.load_page(page_num)
        text = page.get_text()
        results.append({
            "text": text.strip(),
            "metadata": {
                "page": page_num + 1,
                "type": "pdf"
            }
        })
    doc.close()
    return results

def parse_docx(file_path: str):
    """Parses a Word document paragraph by paragraph."""
    doc = Document(file_path)
    full_text = []
    for para in doc.paragraphs:
        if para.text.strip():
            full_text.append(para.text.strip())
            
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                full_text.append(row_text)
                
    text_content = "\n".join(full_text)
    return [{
        "text": text_content,
        "metadata": {
            "page": 1,
            "type": "docx"
        }
    }]

def parse_pptx(file_path: str):
    """Parses a PowerPoint slide by slide."""
    prs = Presentation(file_path)
    results = []
    for i, slide in enumerate(prs.slides):
        slide_text = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                slide_text.append(shape.text.strip())
        results.append({
            "text": "\n".join(slide_text),
            "metadata": {
                "page": i + 1,
                "type": "pptx"
            }
        })
    return results

def parse_excel(file_path: str):
    """Parses Excel spreadsheets sheet by sheet into CSV format."""
    xl = pd.ExcelFile(file_path)
    results = []
    for sheet_name in xl.sheet_names:
        df = xl.parse(sheet_name)
        csv_text = df.to_csv(index=False)
        results.append({
            "text": f"Sheet: {sheet_name}\n\n{csv_text}",
            "metadata": {
                "page": sheet_name,
                "type": "xlsx"
            }
        })
    return results

def parse_image(file_path: str):
    """Parses images using pytesseract OCR if available, else returns placeholder."""
    try:
        import pytesseract
        img = Image.open(file_path)
        text = pytesseract.image_to_string(img)
        return [{
            "text": text.strip(),
            "metadata": {
                "page": 1,
                "type": "image"
            }
        }]
    except Exception as e:
        return [{
            "text": f"[Image uploaded: {os.path.basename(file_path)} - OCR extraction failed or Tesseract is not installed. Path: {file_path}]",
            "metadata": {
                "page": 1,
                "type": "image_error"
            }
        }]

def parse_document(file_path: str):
    """Detects file type and parses using the appropriate parser."""
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".pdf":
        return parse_pdf(file_path)
    elif ext in [".docx", ".doc"]:
        return parse_docx(file_path)
    elif ext in [".pptx", ".ppt"]:
        return parse_pptx(file_path)
    elif ext in [".xlsx", ".xls"]:
        return parse_excel(file_path)
    elif ext in [".png", ".jpg", ".jpeg", ".tiff", ".bmp", ".webp"]:
        return parse_image(file_path)
    elif ext in [".txt", ".csv", ".md"]:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            content = f.read()
        return [{
            "text": content,
            "metadata": {
                "page": 1,
                "type": ext[1:]
            }
        }]
    else:
        raise ValueError(f"Unsupported file extension: {ext}")
